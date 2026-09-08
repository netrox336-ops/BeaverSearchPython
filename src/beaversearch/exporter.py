from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font


HEADERS_RU = ["№", "Ник", "SteamID64", "Steam", "CS2 ₽", "Dota 2 ₽", "Rust ₽", "Итого ₽", "Совпадения", "Сервер", "Карта", "Последняя проверка"]
HEADERS_EN = ["#", "Nickname", "SteamID64", "Steam", "CS2 ₽", "Dota 2 ₽", "Rust ₽", "Total ₽", "Matched", "Server", "Map", "Last check"]


def _row(index: int, r: dict) -> list:
    profile = r.get("profile_url") or f"https://steamcommunity.com/profiles/{r['steam_id64']}"
    return [
        index, r.get("nickname", ""), r.get("steam_id64", ""), profile,
        r.get("cs2_rub"), r.get("dota2_rub"), r.get("rust_rub"), r.get("total_rub"),
        r.get("matched_games", ""), r.get("server_address", ""), r.get("map_name", ""), r.get("last_seen_at", ""),
    ]


def export_xlsx(path: str | Path, rows: list[dict], language: str = "ru") -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    headers = HEADERS_RU if language == "ru" else HEADERS_EN
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, row in enumerate(rows, 1):
        ws.append(_row(i, row))
    widths = [6, 24, 20, 46, 14, 14, 14, 14, 24, 24, 18, 24]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def export_csv(path: str | Path, rows: list[dict], language: str = "ru") -> Path:
    path = Path(path)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(HEADERS_RU if language == "ru" else HEADERS_EN)
        for i, row in enumerate(rows, 1):
            writer.writerow(_row(i, row))
    return path
